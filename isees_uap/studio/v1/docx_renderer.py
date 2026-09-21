"""Deterministic, non-networked DOCX renderer for Studio V1 render models."""
from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from .render_model import RenderDocument, RenderModelFailure, RenderModelFailureCode

RENDERER_VERSION = "studio-v1-python-docx/1"
TEMPLATE_VERSION = "investigation-report-docx/1"
MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
CONFIGURATION_HASH = "sha256:d71b44bcde4fb6847d842df974368a4469dac8783280eefd55f78bb2fd9f1f49"
_FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
_BAD_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_PROHIBITED = ("vbaproject", "activex/", "embeddings/", "oleobject", "externallink")


def _safe(value: str) -> str:
    if _BAD_XML.search(value):
        raise RenderModelFailure(RenderModelFailureCode.INVALID_CHARACTER,
                                 "Document content contains an invalid XML control character.")
    return value


def _paragraph(document: Document, text: str, style: str | None = None):
    paragraph = document.add_paragraph(style=style)
    paragraph.add_run(_safe(text))
    return paragraph


def _normalize_and_validate(raw: bytes) -> bytes:
    source, output = BytesIO(raw), BytesIO()
    with ZipFile(source, "r") as package:
        names = package.namelist()
        lowered = tuple(name.lower() for name in names)
        if any(any(marker in name for marker in _PROHIBITED) for name in lowered):
            raise RenderModelFailure(RenderModelFailureCode.UNSUPPORTED_CONTENT,
                                     "Generated DOCX contains prohibited active content.")
        for name in names:
            if name.endswith(".rels"):
                relationship = package.read(name)
                if b'TargetMode="External"' in relationship or b"TargetMode='External'" in relationship:
                    raise RenderModelFailure(RenderModelFailureCode.UNSUPPORTED_CONTENT,
                                             "Generated DOCX contains an external relationship.")
        required = {"[Content_Types].xml", "_rels/.rels", "word/document.xml"}
        if not required.issubset(names):
            raise RenderModelFailure(RenderModelFailureCode.MALFORMED_CONTENT,
                                     "Generated DOCX package is incomplete.")
        with ZipFile(output, "w", ZIP_DEFLATED, compresslevel=9) as normalized:
            for name in sorted(names):
                info = ZipInfo(name, _FIXED_ZIP_TIME)
                info.compress_type = ZIP_DEFLATED
                info.external_attr = 0o600 << 16
                info.create_system = 3
                normalized.writestr(info, package.read(name), compress_type=ZIP_DEFLATED, compresslevel=9)
    return output.getvalue()


def render_docx(model: RenderDocument, exported_at: str) -> bytes:
    when = datetime.fromisoformat(exported_at.replace("Z", "+00:00")).astimezone(timezone.utc)
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.75)
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    properties = document.core_properties
    properties.title = _safe(model.title)
    properties.author = _safe(model.author_principal_id or "iSEES Studio guest")
    properties.subject = (f"Studio V1 immutable revision {model.revision_id}" if model.source_kind == "SAVED_REVISION"
                          else "Studio V1 LOCAL DRAFT — CURRENT DRAFT — Not an authority record")
    properties.creator = RENDERER_VERSION
    properties.last_modified_by = RENDERER_VERSION
    properties.created = when
    properties.modified = when
    title = _paragraph(document, model.title, "Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if model.source_kind == "CURRENT_DRAFT":
        _paragraph(document, "LOCAL DRAFT", "Heading 1")
        _paragraph(document, "CURRENT DRAFT", "Heading 2")
        _paragraph(document, "Not an authority record")
    else:
        _paragraph(document, "GOVERNED SAVED REVISION", "Heading 1")
        _paragraph(document, "Projection of an immutable .author revision; not an authority record.")
    for block in model.blocks:
        style = None
        if block.kind == "HEADING": style = f"Heading {min(block.level or 1, 6)}"
        elif block.kind in {"SECTION", "APPENDIX"}: style = "Heading 2"
        _paragraph(document, block.text, style)
        if block.placeholder_reason:
            _paragraph(document, f"[{block.kind} placeholder: {block.placeholder_reason}]")
        if block.provenance:
            _paragraph(document, "Provenance: " + ", ".join(block.provenance))
        if block.kind == "TABLE":
            table = document.add_table(rows=1, cols=len(block.columns))
            table.style = "Table Grid"
            for cell, value in zip(table.rows[0].cells, block.columns): cell.text = _safe(value)
            for values in block.rows:
                for cell, value in zip(table.add_row().cells, values): cell.text = _safe(value)
    _paragraph(document, "Provenance", "Heading 1")
    if model.source_kind == "SAVED_REVISION":
        _paragraph(document, f"Artifact: {model.artifact_id}")
        _paragraph(document, f"Revision: {model.revision_id} (r{model.revision_number:04d})")
    else:
        _paragraph(document, f"Document: {model.document_id}")
        _paragraph(document, "Authority: LOCAL DRAFT — CURRENT DRAFT — Not an authority record")
    if model.investigation_id: _paragraph(document, f"Investigation: {model.investigation_id}")
    _paragraph(document, f"Source hash: {model.content_hash}")
    _paragraph(document, f"Profile: {model.profile} · {model.profile_version}")
    _paragraph(document, f"Exported: {exported_at} · Renderer: {RENDERER_VERSION}")
    for snapshot_id, snapshot_hash in model.snapshot_references:
        _paragraph(document, f"Frozen snapshot: {snapshot_id} ({snapshot_hash})")
    if model.citations:
        _paragraph(document, "Citations", "Heading 1")
        for citation in model.citations:
            bits = [", ".join(citation.authors), citation.title, citation.issued or "",
                    citation.publisher or "", citation.url or "", f"snapshot {citation.source_snapshot_id}"]
            _paragraph(document, ". ".join(value for value in bits if value))
    buffer = BytesIO()
    document.save(buffer)
    return _normalize_and_validate(buffer.getvalue())
