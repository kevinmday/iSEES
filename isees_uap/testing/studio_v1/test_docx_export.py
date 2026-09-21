from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from fastapi.testclient import TestClient

from isees_uap.studio.v1.docx_renderer import (CONFIGURATION_HASH, MEDIA_TYPE,
    RENDERER_VERSION, TEMPLATE_VERSION, render_docx)
from isees_uap.studio.v1.hashing import canonical_sha256
from isees_uap.studio.v1.render_model import build_render_model
from isees_uap.testing.studio_v1.test_author_api import api
from isees_uap.testing.studio_v1.test_application import application
from isees_uap.testing.studio_v1.test_current_draft_export import _counts, _wire
from isees_uap.testing.studio_v1.test_persistence import make_command


def _text(data):
    document = Document(BytesIO(data))
    return "\n".join([p.text for p in document.paragraphs] +
                     [cell.text for table in document.tables for row in table.rows for cell in row.cells])


def _create(app, command, key="docx-1"):
    return app.create_export("principal-1", "investigation-1", command.artifact.artifactId,
        command.revision.revisionId, format="DOCX", template_version=TEMPLATE_VERSION,
        renderer_version=RENDERER_VERSION, configuration_hash=CONFIGURATION_HASH,
        idempotency_key=key)


def test_direct_render_is_deterministic_valid_safe_and_unicode_faithful():
    command = make_command(projections=0)
    raw = deepcopy(command.revision.model_dump(exclude_none=True))
    raw["semanticContent"]["title"] = "Relation ↔ “signal” – test — α ≤ β"
    raw["semanticContent"]["nodes"][0]["text"] = "Unicode ΓΔΘ ≠ ± and hash " + "a" * 64
    raw["contentHash"] = canonical_sha256(raw["semanticContent"])
    command = make_command(projections=0, revision=raw)
    model = build_render_model(command.revision, "investigation-1")
    first = render_docx(model, "2026-09-21T18:00:00.000Z")
    assert first == render_docx(model, "2026-09-21T18:00:00.000Z")
    assert first.startswith(b"PK\x03\x04")
    assert "↔" in _text(first) and "ΓΔΘ ≠ ±" in _text(first)
    with ZipFile(BytesIO(first)) as package:
        names = set(package.namelist())
        assert {"[Content_Types].xml", "_rels/.rels", "word/document.xml"} <= names
        assert all(info.date_time == (1980, 1, 1, 0, 0, 0) for info in package.infolist())
        assert not any(any(marker in name.lower() for marker in
            ("vbaproject", "activex/", "embeddings/", "oleobject")) for name in names)
        assert all(b'TargetMode="External"' not in package.read(name)
                   for name in names if name.endswith(".rels"))


def test_guest_current_draft_docx_creates_no_state_and_has_disclaimers(api):
    application_api, _, _, tmp_path = api
    payload = _wire(make_command(projections=0), templateProfileVersion=TEMPLATE_VERSION,
                    rendererVersion=RENDERER_VERSION, configurationHash=CONFIGURATION_HASH)
    before_payload, before = deepcopy(payload), _counts(tmp_path / "studio-v1.sqlite3")
    url = "/api/v1/investigations/investigation-1/studio-v1/artifacts/current-draft/exports/docx"
    with TestClient(application_api) as guest:
        response = guest.post(url, json=payload, headers={"X-ISEES-Current-Draft": "1"})
    assert response.status_code == 200 and response.headers["content-type"] == MEDIA_TYPE
    assert response.headers["content-disposition"].endswith('.docx"')
    assert response.headers["cache-control"] == "private, no-store"
    assert response.headers["pragma"] == "no-cache"
    assert response.headers["x-content-type-options"] == "nosniff"
    text = _text(response.content)
    assert all(value in text for value in ("LOCAL DRAFT", "CURRENT DRAFT", "Not an authority record",
                                            payload["sourceHash"]))
    assert payload == before_payload and _counts(tmp_path / "studio-v1.sqlite3") == before


def test_saved_docx_is_independent_from_pdf_and_uses_immutable_revision(tmp_path):
    app = application(tmp_path); app.start()
    command = make_command(projections=0); app.save("principal-1", "investigation-1", command)
    pdf = app.create_export("principal-1", "investigation-1", command.artifact.artifactId,
        command.revision.revisionId, format="PDF", template_version="investigation-report-pdf/1",
        renderer_version="studio-v1-reportlab-pdf/1",
        configuration_hash="sha256:6d1e0783d1fe839271f281ee34b7355a618291c6e31b1282cfc170e297dab200",
        idempotency_key="pdf-independent")
    docx = _create(app, command)
    record, data = app.download_export("principal-1", "investigation-1", command.artifact.artifactId,
                                      command.revision.revisionId, docx.export_id)
    jobs = app.list_projection_jobs("principal-1", "investigation-1", command.artifact.artifactId,
                                    command.revision.revisionId)
    assert record.media_type == MEDIA_TYPE and record.safe_filename.endswith(".docx")
    assert docx.export_id != pdf.export_id and docx.output_hash != pdf.output_hash
    assert {job.format for job in jobs} >= {"PDF", "DOCX"}
    text = _text(data)
    assert command.revision.revisionId in text and command.revision.contentHash in text
    assert command.revision.semanticContent.title in text
    app.close()
